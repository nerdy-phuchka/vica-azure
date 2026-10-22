"""Optional: produce DOCX and XLSX fixtures from the supplied synthetic UTF-8 files."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
import csv
root=Path(__file__).resolve().parent
for filename in ('01_registration_bilingual.txt','02_financial_statements_bilingual.txt'):
    doc=Document()
    for line in (root/filename).read_text(encoding='utf-8').splitlines():
        doc.add_paragraph(line)
    doc.save(root/filename.replace('.txt','.docx'))
book=Workbook();sheet=book.active;sheet.title='Synthetic statement'
for row in csv.reader((root/'03_bank_statement_bilingual.csv').read_text(encoding='utf-8').splitlines()):sheet.append(row)
book.save(root/'03_bank_statement_bilingual.xlsx')
print('Generated synthetic Office samples only.')
