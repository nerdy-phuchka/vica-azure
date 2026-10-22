import html
import io
from docx import Document


def report_lines(case,assessment,kind='credit_report'):
    data=case.body()
    lines=[('Report',kind.replace('_',' ').title()),('Application',case.id),
      ('Borrower',data['borrower_name']),('Status',data['status']),
      ('Requested amount',f"{data['currency']} {data['requested_amount']}"),
      ('Product',data['product']),('Purpose',data['purpose']),
      ('Proposed tenor (months)',str(data.get('tenor_months',12))),
      ('Proposed security',data.get('security_details','Not provided')),
      ('Classification','SYNTHETIC DEMO' if data['synthetic'] else 'Confidential — real borrower information')]
    if not assessment:
        return lines+[('Assessment','Not yet generated')]
    a=assessment.body()
    lines += [('Assessment ID',assessment.id),('Reporting period',a['period']),
              ('Policy version',a['policy']['version']),('Recommendation',a['recommendation']),('Rating',a['rating'])]
    if kind=='field_visit':
        return lines+[(n['category'],n['text']) for n in a['notes'] if n['category']=='field_visit']
    lines += [('Outstanding issue',b) for b in a['blockers']]
    lines += [('Metric: '+k, str(v) if v is not None else 'Unavailable') for k,v in a['financial']['metrics'].items()]
    lines += [('Assumption',s) for s in a['financial']['assumptions']]
    lines += [(c['rule_id'],c['description']+' — '+c['outcome']) for c in a['checks']]
    if kind=='credit_report':
        lines += [('Evidence: '+f['name'],f"{f['value']} | Document {f.get('document_id')} | Page {f.get('page')} | {f.get('quote','')} | Reviewer {f.get('reviewed_by')}") for f in a['fields']]
        lines += [('Review: '+k,str(v)) for k,v in a['clearances'].items()]
        lines += [(n['category'],n['text']) for n in a['notes']]
    if data.get('decision'):
        lines += [('Decision',str(data['decision']))]
    return lines


def render_html(lines,language='en'):
    esc=html.escape
    direction='rtl' if language=='ar' else 'ltr'
    rows=''.join(f'<tr><th>{esc(k)}</th><td>{esc(v)}</td></tr>' for k,v in lines)
    return f'''<!doctype html><html lang="{language}" dir="{direction}"><meta charset="utf-8">
<title>Virtual Credit Analyst — report</title><style>body{{font:15px Arial;margin:40px;color:#163047}}
table{{width:100%;border-collapse:collapse}}th,td{{padding:12px;border-bottom:1px solid #ddd;text-align:start;white-space:pre-wrap;overflow-wrap:anywhere}}th{{width:25%}}h1{{color:#146b70}}</style>
<h1>Virtual Credit Analyst / محلل الائتمان الافتراضي</h1><p>Human-reviewed prototype assessment. No automatic lending decision.</p><table>{rows}</table></html>'''


def render_docx(lines):
    doc=Document()
    doc.add_heading('Virtual Credit Analyst / محلل الائتمان الافتراضي',0)
    for title,text in lines:
        doc.add_heading(title,level=2)
        doc.add_paragraph(text)
    output=io.BytesIO()
    doc.save(output)
    return output.getvalue()
