"""Strict, human-reviewed statement schedules. OCR statements are available for manual transcription."""
import csv
import io
from .finance import number,fmt

CATEGORIES={'operating_receipt','operating_payment','loan_proceeds','owner_injection',
            'internal_transfer','debt_service','other','unclassified'}
ALIASES={'date':{'date','التاريخ'},'description':{'description','الوصف','البيان'},
         'debit':{'debit','مدين'},'credit':{'credit','دائن'},'balance':{'balance','الرصيد'}}


def parse_schedule(filename,content):
    if filename.lower().endswith('.csv'):
        rows=list(csv.reader(io.StringIO(content.decode('utf-8-sig'))))
    elif filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        sheet=load_workbook(io.BytesIO(content),read_only=True,data_only=True).active
        if sheet.max_row>20000 or sheet.max_column>100:
            raise ValueError('Statement schedule exceeds dimensions')
        rows=list(sheet.iter_rows(values_only=True))
    else:
        return []
    if not rows:return []
    headers=[str(x).strip().lower() for x in rows[0]]
    mapping={key:next((i for i,h in enumerate(headers) if h in aliases),None) for key,aliases in ALIASES.items()}
    if any(mapping[k] is None for k in ('date','description','debit','credit')):
        return []  # Not a transaction schedule; leave as generic extracted text.
    transactions=[]
    for index,row in enumerate(rows[1:],2):
        if not any(x is not None and str(x).strip() for x in row):continue
        def cell(key):
            i=mapping[key]
            return str(row[i]) if i is not None and i<len(row) and row[i] is not None else ''
        debit=number(cell('debit')) or number('0')
        credit=number(cell('credit')) or number('0')
        if debit<0 or credit<0 or (debit>0 and credit>0):
            raise ValueError(f'Invalid debit/credit amounts at statement row {index}')
        transactions.append({'row':index,'date':cell('date'),'description':cell('description'),
           'debit':str(debit),'credit':str(credit),'balance':fmt(number(cell('balance'))),
           'category':'unclassified','reviewed':False})
    return transactions


def summarise(transactions):
    totals={key:number('0') for key in CATEGORIES}
    unreviewed=0
    for tx in transactions:
        if not tx.get('reviewed'):
            unreviewed+=1;continue
        totals[tx['category']]+=number(tx['credit'])-number(tx['debit'])
    return {'reviewed_net_by_category':{key:fmt(value) for key,value in totals.items()},
            'unreviewed_count':unreviewed,
            'note':'Reviewed net cash movements, not accounting sales. Reconcile all banking relationships separately.'}
